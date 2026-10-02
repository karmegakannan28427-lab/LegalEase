"""
TXT, DOCX and PDF export utilities.
"""

from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF

from backend.utils.text_utils import sanitize_text


def _split_lines(text: str) -> list[str]:
    """
    Split document text into clean lines.
    """

    return [
        line.strip()
        for line in sanitize_text(text).splitlines()
    ]


def format_txt(text: str) -> bytes:
    """
    Convert document into TXT bytes.
    """

    return sanitize_text(text).encode(
        "utf-8"
    )


def format_docx(
    text: str,
    doc_type: str,
    terms: str = "",
    logo_path: str | None = None,
) -> bytes:
    """
    Generate a DOCX document.
    """

    document = Document()


    # Page margins
    section = document.sections[0]

    section.top_margin = Inches(0.7)

    section.bottom_margin = Inches(0.7)

    section.left_margin = Inches(0.8)

    section.right_margin = Inches(0.8)


    # Default font
    styles = document.styles

    normal = styles["Normal"]

    normal.font.name = "Times New Roman"

    normal.font.size = Pt(11)


    # Logo
    if logo_path and Path(logo_path).exists():

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run()

        run.add_picture(
            logo_path,
            width=Inches(1.1)
        )


    # Document title
    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        doc_type.upper()
    )

    run.bold = True

    run.font.name = "Times New Roman"

    run.font.size = Pt(16)


    # Main content
    for line in _split_lines(text):

        if not line:

            document.add_paragraph("")

            continue


        if _is_heading(line):

            paragraph = document.add_paragraph()

            run = paragraph.add_run(line)

            run.bold = True

            run.font.name = "Times New Roman"

            run.font.size = Pt(12)

        else:

            paragraph = document.add_paragraph(
                line
            )

            paragraph.paragraph_format.space_after = Pt(6)

            paragraph.paragraph_format.line_spacing = 1.15


    # Terms table
    clean_terms = [
        term.strip()
        for term in terms.split(";")
        if term.strip()
    ]


    if clean_terms:

        document.add_paragraph("")

        heading = document.add_paragraph()

        heading.add_run(
            "Terms Provided by User"
        ).bold = True


        table = document.add_table(
            rows=1,
            cols=2
        )

        table.style = "Table Grid"


        table.rows[0].cells[0].text = "No."

        table.rows[0].cells[1].text = "Term"


        for index, term in enumerate(
            clean_terms,
            start=1
        ):

            cells = table.add_row().cells

            cells[0].text = str(index)

            cells[1].text = term


    # Footer
    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer.add_run(
        "LegalEase - AI-generated draft. "
        "Review by a qualified legal professional."
    )


    # Save to memory
    buffer = BytesIO()

    document.save(buffer)

    return buffer.getvalue()


def _is_heading(line: str) -> bool:
    """
    Detect headings.
    """

    upper = line.upper()


    # Numbered heading
    if re.match(
        r"^\d+[\.\)]\s+",
        line
    ):

        return True


    # ALL CAPS heading
    if (
        len(line) <= 80
        and upper == line
        and any(
            character.isalpha()
            for character in line
        )
    ):

        return True


    return False


def _pdf_safe(text: str) -> str:
    """
    Convert text into PDF-safe Latin-1 text.
    """

    return (
        sanitize_text(text)
        .replace("\u2022", "-")
        .encode(
            "latin-1",
            "replace"
        )
        .decode("latin-1")
    )


class LegalEasePDF(FPDF):

    def __init__(
        self,
        logo_path: str | None = None
    ):

        super().__init__()

        self.logo_path = logo_path


    def header(self):

        if (
            self.logo_path
            and Path(self.logo_path).exists()
        ):

            self.image(
                self.logo_path,
                x=95,
                y=8,
                w=20
            )

            self.ln(22)

        else:

            self.ln(5)


    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            size=8
        )

        self.cell(
            0,
            10,
            "LegalEase - AI-generated draft | "
            "Review by a qualified legal professional",
            align="C"
        )


def format_pdf(
    text: str,
    doc_type: str,
    logo_path: str | None = None,
) -> bytes:
    """
    Generate PDF document.
    """

    pdf = LegalEasePDF(
        logo_path=logo_path
    )


    pdf.set_auto_page_break(
        auto=True,
        margin=18
    )


    pdf.add_page()


    # Title
    pdf.set_font(
        "Helvetica",
        "B",
        16
    )

    pdf.multi_cell(
        0,
        10,
        _pdf_safe(
            doc_type.upper()
        ),
        align="C"
    )

    pdf.ln(4)


    # Content
    for line in _split_lines(text):

        if not line:

            pdf.ln(4)

            continue


        if _is_heading(line):

            pdf.set_font(
                "Helvetica",
                "B",
                11
            )

            pdf.multi_cell(
                0,
                7,
                _pdf_safe(line)
            )

            pdf.ln(1)

        else:

            pdf.set_font(
                "Helvetica",
                size=10.5
            )

            pdf.multi_cell(
                0,
                6,
                _pdf_safe(line)
            )

            pdf.ln(1)


    output = pdf.output(
        dest="S"
    )


    return bytes(output)